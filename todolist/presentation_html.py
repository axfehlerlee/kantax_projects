from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

OUT=Path('/Users/geil/Downloads/AI_OUTPUTS/2026-10-02_0609_html-presentation-v2')
OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
prs.core_properties.title='내 업보들 — HTML을 활용한 웹 문서 구조 설계'
prs.core_properties.author='Codex'
FONT='Malgun Gothic'
preview_font='/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/malgun.ttf'
sources=[('1장 1강: HTML의 개념과 텍스트 태그','https://app.notion.com/p/536e0c9f801583a9a20f01ba3cca1b63'),('1장 2강: 하이퍼링크, 멀티미디어 및 폼 태그','https://app.notion.com/p/9d1e0c9f8015834eab5c01c63107b31e'),('1장 3강: HTML5 시멘틱 태그의 구조화','https://app.notion.com/p/e8ce0c9f8015829eb3d7810a6a128b8c')]
notes=[
'안녕하세요. 저는 할 일 관리 페이지, 내 업보들을 만들었습니다. 오늘은 1장 수업과 제가 정리한 학습 노트를 바탕으로, 이 페이지의 HTML 구조를 설명하겠습니다. 핵심은 태그를 모양으로 고르는 것이 아니라, 콘텐츠의 역할에 맞춰 선택하는 것입니다.',
'먼저 문서 구조와 텍스트 태그입니다. doctype으로 HTML 문서를 선언하고, lang을 ko로, 인코딩을 UTF-8로 지정했습니다. head의 title은 브라우저 탭 제목이고, body의 h1은 화면 본문의 대표 제목입니다. 둘 다 내 업보들이지만 역할은 다릅니다. 세 상태 영역은 h2로 표현해서 h1 아래의 제목 위계를 만들었습니다. 날짜와 안내는 p, 개수와 카드 제목은 span으로 묶었습니다. 수업과 노트에서 배운 대로 제목 수준은 글자 크기가 아닌 의미로 선택했고, 크기와 배치는 CSS로 조절했습니다.',
'다음은 폼입니다. form 안에 텍스트 input과 submit 버튼을 넣었습니다. id는 JavaScript가 요소를 찾을 때 쓰는 식별자이고, name은 기본 폼 전송에서 데이터의 키가 됩니다. 사용자가 입력한 값은 JavaScript의 input.value로 읽습니다. required와 maxlength는 빈 입력과 입력 길이를 제한합니다. 노트에서 배운 action은 제출 대상이고 method는 전송 방식입니다. 하지만 제 앱은 서버 제출 대상과 전송 방식을 별도로 지정하지 않았습니다. 제출 이벤트에서 preventDefault로 기본 전송을 막고, 입력을 읽어 localStorage에 저장합니다. 그래서 폼 구조는 구현했지만 GET이나 POST로 서버에 저장하는 기능은 구현한 것이 아닙니다.',
'시멘틱 레이아웃은 실제 HTML의 중첩 구조로 설명하겠습니다. main은 핵심 콘텐츠를 감싸고, header에는 제목과 날짜가 들어갑니다. 보드 section 안에는 입력 폼과 세 상태 section, 그리고 보드의 footer가 있습니다. 각 상태 section은 h2 제목과 연결했습니다. footer는 페이지 전체가 아니라 이 보드의 남은 개수와 정리 기능을 묶는 영역입니다. 반면 div는 입력창과 버튼을 나란히 놓기 위한 배치 묶음입니다. ul은 할 일 목록, li는 개별 항목이고, template은 반복할 카드 구조를 보관합니다. JavaScript가 template을 복제해 화면에 넣습니다. 이렇게 구조의 의미와 화면의 모양을 구분했습니다.',
'현재는 할 일 추가, 제목 수정, 삭제와 브라우저 저장을 구현했습니다. In Progress와 Done 칸은 있지만 상태를 변경하는 기능은 아직 구현하지 못했습니다. 새 항목은 Backlog에 들어갑니다. 이번 작업에서는 배운 태그를 모두 넣기보다 필요한 구조에 적용했습니다. 링크나 미디어, strong과 b, del과 s의 구분은 배웠지만 현재 화면에는 사용하지 않았습니다. HTML은 의미와 입력 구조, CSS는 디자인, JavaScript는 동작을 담당한다는 점을 코드로 정리했습니다. 감사합니다.'
]

sources += [('학습 노트: HTML 폼과 데이터 전송','https://app.notion.com/p/3dce0c9f801581ba8eb1eabe9a427d30'), ('학습 노트: HTML 의미 구조와 텍스트 태그','https://app.notion.com/p/3dbe0c9f801581fe895edd1f9ad0604b'), ('WHATWG: 폼 데이터 구성','https://html.spec.whatwg.org/multipage/form-control-infrastructure.html#constructing-the-entry-list'), ('WHATWG: 제목과 영역','https://html.spec.whatwg.org/multipage/sections.html')]
images=[]; text_manifest=[]
def base(num,title,kicker,dark=False):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string('000000' if dark else 'FFFFFF')
    im=Image.new('RGB',(1600,900),'black' if dark else 'white'); draw=ImageDraw.Draw(im)
    state={'s':slide,'im':im,'d':draw,'fg':'FFFFFF' if dark else '000000','texts':[]}
    tx(state,kicker,0.65,0.48,11.9,0.3,12)
    tx(state,title,0.65,1.0,12,0.8,32,True)
    tx(state,f'{num:02d} / 05',11.75,6.95,0.9,0.25,10)
    return state

def tx(st,text,x,y,w,h,size=20,bold=False,color=None):
    color=color or st['fg']; box=st['s'].shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(5)
    font=ImageFont.truetype(preview_font,round(size*1600/960))
    yy=y*120
    for line in text.split('\n'):
        st['d'].text((x*120,yy),line,font=font,fill='#'+color)
        if st['d'].textlength(line,font=font)>w*120: raise ValueError('Preview line overflow: '+line)
        yy+=size*1600/960*1.35
    if yy-y*120>h*120+8: raise ValueError('Preview height overflow: '+text)
    st['texts'].append(text)

def rect(st,x,y,w,h,fill=None):
    fill=fill or ('000000' if st['fg']=='FFFFFF' else 'FFFFFF'); sh=st['s'].shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.color.rgb=RGBColor.from_string(st['fg']); sh.line.width=Pt(1)
    st['d'].rectangle((x*120,y*120,(x+w)*120,(y+h)*120),fill='#'+fill,outline='#'+st['fg'],width=2)

def finish(st,i):
    st['s'].notes_slide.notes_text_frame.text=notes[i]+'\n\n참고 자료\n'+'\n'.join(n+'\n'+u for n,u in sources)+'\n로컬 근거: index.html, js/todo.js (2026-10-02 확인)'
    path=OUT/f'slide-{i+1}.png'; st['im'].save(path); images.append(st['im']); text_manifest.append(st['texts'])

s=base(1,'내 업보들','HTML 시멘틱 레이아웃 · 수업과 학습 노트 적용 · 약 3분',True)
tx(s,'할 일은 쌓이고,\n문서에는 구조가 필요합니다.',0.65,2.3,7,1.6,29,True)
for i,(tag,label) in enumerate([('01','문서와 텍스트'),('02','사용자 입력 폼'),('03','의미 있는 영역')]):
    y=2.35+i*1.18; rect(s,8.5,y,4.05,0.88); tx(s,tag+'  '+label,8.75,y+0.2,3.5,0.45,20)
tx(s,'태그의 기본 모양과 의미를 구분해 설계했습니다.',0.65,5.55,8,0.5,18)
finish(s,0)

s=base(2,'같은 제목, 서로 다른 역할','1강 + 학습 노트 · HTML 의미 구조와 텍스트 태그')
rect(s,0.65,2.2,5.8,3.6)
tx(s,'head  문서 정보',0.95,2.5,5.1,0.45,22,True)
tx(s,'<title>내 업보들</title>\n브라우저 탭의 제목',0.95,3.15,5.1,1.1,20)
tx(s,'문서 선언: doctype / html: lang="ko" / head: UTF-8',0.65,5.95,12,0.4,17)
rect(s,6.85,2.2,5.8,3.6)
tx(s,'body  화면 콘텐츠',7.15,2.5,5.1,0.45,22,True)
tx(s,'h1  내 업보들 → 대표 제목\nh2  세 상태 영역 → 하위 제목\np / span  안내와 텍스트 묶음',7.15,3.15,5.1,1.9,19)
tx(s,'제목 수준은 의미로 선택하고, 크기와 배치는 CSS로 조절합니다.',0.65,6.5,12,0.4,18,True)
finish(s,1)

s=base(3,'폼을 만들었다 ≠ 서버에 전송했다','2강 + 학습 노트 · HTML 폼과 데이터 전송')
for x,title,body in [(0.65,'id="todo-input"','요소를 찾는 식별자'),(4.8,'name="todo"','기본 폼 전송의 키'),(8.95,'input.value','현재 입력된 값 읽기')]:
    rect(s,x,2.25,3.7,1.4);tx(s,title,x+0.2,2.48,3.3,0.45,19,True);tx(s,body,x+0.2,3.08,3.3,0.35,16)
tx(s,'form 안에 input + button(type="submit")\nrequired: 필수 입력  /  maxlength="200": 입력 길이 제한',0.65,4.0,12,1.05,20)
tx(s,'제출 → preventDefault() → input.value → localStorage',0.65,5.5,12,0.5,21,True)
tx(s,'action·method 별도 지정 없음 / GET·POST 서버 저장 미구현',0.65,6.28,12,0.4,17)
finish(s,2)

s=base(4,'의미 있는 구역으로 나눴습니다','3강 · HTML5 시멘틱 태그의 구조화')
rect(s,0.65,2.1,6.25,4.65); tx(s,'main  핵심 콘텐츠',0.9,2.28,5.7,0.4,19,True)
rect(s,0.95,2.9,5.65,0.6); tx(s,'header  제목 · 날짜',1.15,3.0,5.2,0.35,17)
rect(s,0.95,3.8,5.65,2.65); tx(s,'section  할 일 보드',1.15,3.96,5.2,0.35,17,True)
tx(s,'form  입력창 · 추가 버튼',1.3,4.52,4.9,0.35,16)
tx(s,'section  Backlog / In Progress / Done',1.3,5.05,4.9,0.35,16)
rect(s,1.2,5.67,5.1,0.55); tx(s,'footer  보드의 개수 · 정리',1.4,5.77,4.7,0.35,16)
tx(s,'section + h2\n각 상태 영역의 주제를 표시',7.45,2.3,5.2,1.0,20)
tx(s,'div\n배치를 위한 묶음',7.45,3.65,5.2,1.0,20)
tx(s,'ul / li + template\n목록과 반복할 카드 구조',7.45,5.0,5.2,1.0,20)
finish(s,3)

s=base(5,'화면 구조와 구현 기능을 구분합니다','구현 범위 · 배운 점',True)
rect(s,0.65,2.3,5.8,2.4); tx(s,'구현한 기능',0.95,2.58,5.1,0.45,23,True)
tx(s,'추가 · 제목 수정 · 삭제\n브라우저 저장: localStorage',0.95,3.35,5.1,1.1,20)
rect(s,6.85,2.3,5.8,2.4); tx(s,'아직 구현하지 못한 기능',7.15,2.58,5.1,0.45,23,True)
tx(s,'상태 변경: In Progress / Done\n새 항목은 Backlog에 추가',7.15,3.35,5.1,1.1,20)
tx(s,'HTML: 의미 · CSS: 모양 · JavaScript: 동작',0.65,5.55,12,0.65,25,True)
finish(s,4)
prs.save(OUT/'내_업보들_HTML_3분발표_v2.pptx')
images[0].save(OUT/'슬라이드_미리보기.pdf',save_all=True,append_images=images[1:],resolution=120)
thumb=Image.new('RGB',(800,450*5),'white')
for i,im in enumerate(images): thumb.paste(im.resize((800,450)),(0,450*i))
thumb.save(OUT/'전체_미리보기.png')
durations=[15,35,50,50,30]
md='# 내 업보들 — HTML 웹 문서 구조 설계\n\n발표 목표: 약 3분 / 5장 / 블랙앤화이트 / PPTX 글꼴: 맑은 고딕(Malgun Gothic)\n\n'
for i,n in enumerate(notes):md+=f'## {i+1}장 · 약 {durations[i]}초\n\n{n}\n\n'
md+='## 발표 전 확인\n\n- 시간 배분은 계획값입니다. 직접 읽으며 속도를 조절하세요.\n- PPTX 글꼴은 맑은 고딕으로 지정했습니다. 글꼴 파일은 포함하지 않았습니다.\n- PDF·PNG는 동일 배치로 만든 보조 미리보기이며 맑은 고딕을 사용합니다. PowerPoint 실제 렌더링 결과는 아닙니다.\n- 상태 변경은 미구현입니다. 드롭다운이나 카드 이동을 시연하지 마세요.\n- 링크·이미지·영상 태그를 구현했다고 말하지 마세요.\n- 날짜·개수 표시와 추가·수정·삭제·저장은 JavaScript가 담당합니다.\n\n## 학습 내용과 실제 적용\n\n- 적용: title/h1 역할 구분, h1/h2 위계, 시멘틱 영역, 입력 폼 속성.\n- 구분: name은 기본 폼 전송의 키. 현재 JavaScript는 id로 요소를 찾아 .value를 읽으므로 name을 제거해도 이 코드의 값 읽기는 가능하지만, 기본 폼 데이터 구성에서 해당 입력 값은 제외됩니다.\n- 현재 입력창에는 label 요소 대신 aria-label로 이름을 지정했습니다. placeholder가 label을 대신한다고 설명하지 않습니다.\n- 미사용: strong/b, del/s, 체크박스/라디오, readonly, 링크·미디어. 미사용 개념은 구현 성과로 말하지 않습니다.\n\n## 질문 대비\n\n- GET/POST의 차이? GET 폼은 URL 쿼리에, POST 폼은 요청 본문에 데이터를 담습니다. 현재 앱의 할 일 저장에는 둘 다 사용하지 않습니다.\n- POST라서 안전한가? 그 자체로 암호화되지 않습니다. 전송 보호에는 HTTPS가 필요합니다.\n- readonly/disabled? 텍스트 입력의 readonly는 보통 제출에 포함되고, disabled는 제외됩니다. 현재 입력창에는 둘 다 없습니다.\n- value 없는 체크박스? 선택된 경우 기본 값 on이 제출됩니다. 현재 앱에는 체크박스가 없습니다.\n- strong/b? strong은 중요성, b는 특별한 중요성 없이 주의를 끄는 텍스트에 사용합니다. 굵은 모양만 보고 고르지 않습니다.\n- del/s? del은 문서에서 삭제한 내용, s는 더 이상 정확하거나 관련 없는 내용입니다. strike는 폐기된 요소입니다.\n\n## 참고 자료\n\n'
for n,u in sources:md+=f'- [{n}]({u})\n'
md+='\n로컬 근거: /Users/geil/Downloads/jungeun/index.html, js/todo.js\n확인일: 2026-10-02 KST\n'
(OUT/'발표대본.md').write_text(md)
(OUT/'검증기록.md').write_text('# 제작·검증 기록\n\n- 2026-10-02 06:09 KST / #execute #record (의도 추론)\n- 요청: 현재 HTML과 세 강의 및 사용자 학습 노트 두 편을 연결한 3분 발표자료 v2.\n- 작업자: Codex / python-pptx, Pillow, Notion fetch.\n- 강의 세 페이지, 추가 학습 노트 두 편 및 WHATWG 근거와 현재 HTML을 확인. 강의 전체를 구현했다고 주장하지 않음.\n- 디자인: 검정·흰색, PPTX 모든 텍스트 Malgun Gothic 지정.\n- 원본 index.html과 js/todo.js는 변경하지 않음.\n- 글꼴 포함 및 실제 PowerPoint 렌더링은 미검증.\n')
print(OUT)
